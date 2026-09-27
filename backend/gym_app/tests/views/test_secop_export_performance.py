"""Performance and workbook contracts for the SECOP export endpoint."""
from datetime import date, datetime, timedelta
from datetime import timezone as dt_timezone
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework import status

from gym_app.models import SECOPProcess

User = get_user_model()
MAX_SECOP_EXPORT_QUERIES = 1
EXPORT_HEADERS = [
    "Referencia",
    "Entidad",
    "Departamento",
    "Objeto",
    "Modalidad",
    "Tipo Contrato",
    "Estado",
    "Presupuesto",
    "Fecha Publicación",
    "Fecha Cierre",
    "URL",
]
FILTERED_EXPORT_ROW = (
    "REF-FILTERED",
    "Filtered entity",
    "Antioquia",
    "A" * 200,
    "Directa",
    "Servicios",
    "Abierto",
    None,
    None,
    None,
    "https://example.com/filtered",
)


def _create_lawyer():
    """Create the authenticated role used by SECOP endpoint contracts."""
    return User.objects.create_user(
        email="secop-export-lawyer@example.com",
        password=None,
        role="lawyer",
        is_gym_lawyer=True,
    )


def _workbook_rows(response):
    """Read the actual XLSX payload returned to the API consumer."""
    workbook = load_workbook(BytesIO(response.content))
    return list(workbook.active.iter_rows(values_only=True))


def _create_export_process(index, publication_date):
    """Create a distinct process for ordering and export limits."""
    return SECOPProcess.objects.create(
        process_id=f"EXPORT-{index:03d}",
        reference=f"REF-{index:03d}",
        entity_name=f"Export entity {index}",
        department="Cundinamarca",
        procedure_name=f"Procedure {index}",
        publication_date=publication_date,
    )


def _create_export_processes(start, count):
    """Create a consecutive export dataset outside an endpoint test body."""
    return [
        _create_export_process(index, date(2026, 1, 1) + timedelta(days=index))
        for index in range(start, start + count)
    ]


def _create_filtered_processes():
    """Create included and excluded rows for the filtered workbook contract."""
    SECOPProcess.objects.create(
        process_id="EXPORT-FILTERED",
        reference="REF-FILTERED",
        entity_name="Filtered entity",
        department="Antioquia",
        procedure_name="A" * 201,
        procurement_method="Directa",
        contract_type="Servicios",
        status="Abierto",
        process_url="https://example.com/filtered",
        description="needle-in-description",
        raw_data={"large": "payload"},
    )
    SECOPProcess.objects.create(
        process_id="EXPORT-EXCLUDED",
        reference="REF-EXCLUDED",
        entity_name="Excluded entity",
        department="Bogotá",
        procedure_name="Excluded procedure",
        description="other-description",
    )


@pytest.mark.django_db
def test_secop_export_returns_filtered_workbook_with_limited_projection(api_client):
    """Fails if export reads raw fields or stops preserving the eleven-column workbook contract."""
    lawyer = _create_lawyer()
    _create_filtered_processes()
    api_client.force_authenticate(user=lawyer)

    with CaptureQueriesContext(connection) as queries:
        response = api_client.get(
            reverse("secop-export-excel"),
            {"department": "Antioquia", "search": "needle-in-description"},
        )
    data_query = queries[0]["sql"]
    select_clause = data_query.lower().partition(" from ")[0]
    rows = _workbook_rows(response)

    assert response.status_code == status.HTTP_200_OK
    assert len(queries) == MAX_SECOP_EXPORT_QUERIES
    assert "description" not in select_clause
    assert "raw_data" not in select_clause
    assert rows == [
        tuple(EXPORT_HEADERS),
        FILTERED_EXPORT_ROW,
    ]


@pytest.mark.django_db
def test_secop_export_keeps_one_query_for_fifty_rows(api_client):
    """Fails if export grows from one query when the selected row count reaches fifty."""
    lawyer = _create_lawyer()
    _create_export_process(0, date(2026, 1, 1))
    api_client.force_authenticate(user=lawyer)
    url = reverse("secop-export-excel")

    with CaptureQueriesContext(connection) as one_row_queries:
        one_row_response = api_client.get(url)
    _create_export_processes(1, 49)
    with CaptureQueriesContext(connection) as fifty_row_queries:
        fifty_row_response = api_client.get(url)

    assert one_row_response.status_code == status.HTTP_200_OK
    assert len(_workbook_rows(one_row_response)) == 2
    assert fifty_row_response.status_code == status.HTTP_200_OK
    assert len(_workbook_rows(fifty_row_response)) == 51
    assert len(one_row_queries) == len(fifty_row_queries)
    assert len(fifty_row_queries) == MAX_SECOP_EXPORT_QUERIES


@pytest.mark.django_db
def test_secop_export_serializes_zero_price_as_empty_cell(api_client):
    """Fails if a zero price stops matching the blank-value export contract."""
    SECOPProcess.objects.create(
        process_id="EXPORT-ZERO",
        reference="REF-ZERO",
        entity_name="Zero price entity",
        department="Cundinamarca",
        procedure_name="Zero price process",
        base_price=0,
    )
    api_client.force_authenticate(user=_create_lawyer())

    response = api_client.get(reverse("secop-export-excel"))
    rows = _workbook_rows(response)

    assert response.status_code == status.HTTP_200_OK
    assert rows[1][7] is None


@pytest.mark.django_db
def test_secop_export_serializes_positive_price_with_dates(api_client):
    """Fails if an export loses concrete financial values or formatted publication dates."""
    SECOPProcess.objects.create(
        process_id="EXPORT-POSITIVE",
        reference="REF-POSITIVE",
        entity_name="Positive price entity",
        department="Cundinamarca",
        procedure_name="Positive price process",
        base_price=1234.50,
        publication_date=date(2026, 2, 3),
        closing_date=datetime(2026, 2, 4, 10, 15, tzinfo=dt_timezone.utc),
    )
    api_client.force_authenticate(user=_create_lawyer())

    response = api_client.get(reverse("secop-export-excel"))
    rows = _workbook_rows(response)

    assert response.status_code == status.HTTP_200_OK
    assert rows[1][7] == 1234.5
    assert rows[1][8] == "2026-02-03"
    assert rows[1][9] == "2026-02-04 10:15:00+00:00"


@pytest.mark.django_db
def test_secop_export_keeps_newest_five_hundred_rows(api_client):
    """Fails if export exceeds its row limit or stops ordering by newest publication date."""
    lawyer = _create_lawyer()
    _create_export_processes(0, 501)
    api_client.force_authenticate(user=lawyer)

    response = api_client.get(reverse("secop-export-excel"))
    rows = _workbook_rows(response)

    assert response.status_code == status.HTTP_200_OK
    assert len(rows) == 501
    assert rows[1][0] == "REF-500"
    assert rows[-1][0] == "REF-001"


@pytest.mark.django_db
def test_secop_export_returns_headers_for_empty_result(api_client):
    """Fails if an empty filter result produces an invalid workbook without export headers."""
    api_client.force_authenticate(user=_create_lawyer())

    response = api_client.get(reverse("secop-export-excel"), {"department": "No rows"})

    assert response.status_code == status.HTTP_200_OK
    assert _workbook_rows(response) == [tuple(EXPORT_HEADERS)]


@pytest.mark.django_db
def test_secop_export_requires_authentication(api_client):
    """Fails if an anonymous user can download the SECOP export."""
    response = api_client.get(reverse("secop-export-excel"))

    assert response.status_code == status.HTTP_401_UNAUTHORIZED
