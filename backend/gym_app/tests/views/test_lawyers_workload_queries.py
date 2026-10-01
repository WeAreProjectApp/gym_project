"""Query-budget regressions for the lawyers workload Excel export."""

import datetime
import io

import pandas as pd
import pytest
from django.db import connection
from django.http import HttpResponse
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from gym_app.models import Case, Process, Stage, User
from gym_app.views.reports import generate_lawyers_workload_report
from gym_app.views.reports.user_reports import MAX_LAWYERS_WORKLOAD_QUERIES

pytestmark = pytest.mark.django_db

REPORT_START = datetime.datetime(2026, 1, 1, tzinfo=datetime.timezone.utc)
REPORT_END = datetime.datetime(2026, 1, 31, 23, 59, tzinfo=datetime.timezone.utc)


def _response():
    return HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def _lawyer(number):
    return User.objects.create_user(
        email=f"workload-lawyer-{number}@test.com",
        password="testpassword",
        first_name="Workload",
        last_name=f"Lawyer {number}",
        role="lawyer",
        is_gym_lawyer=True,
    )


def _process(lawyer, case, reference, created_at, stages=()):
    process = Process.objects.create(
        authority="Court",
        plaintiff="Plaintiff",
        defendant="Defendant",
        ref=reference,
        lawyer=lawyer,
        case=case,
        subcase="Workload",
    )
    Process.objects.filter(pk=process.pk).update(created_at=created_at)
    process.stages.add(*stages)
    return process


def _workload_query_count():
    with CaptureQueriesContext(connection) as queries:
        generate_lawyers_workload_report(_response(), REPORT_START, REPORT_END)
    return len(queries)


def test_workload_export_keeps_query_count_constant_from_one_to_fifty_lawyers(
    monkeypatch,
):
    """Fails if workload aggregation returns to querying processes or stages per lawyer."""
    monkeypatch.setattr("gym_app.views.reports.user_id", None)
    civil = Case.objects.create(type="Civil")
    _process(_lawyer(1), civil, "ONE", REPORT_START)
    one_lawyer_queries = _workload_query_count()

    Process.objects.all().delete()
    User.objects.filter(role="lawyer").delete()
    for number in range(50):
        lawyer = _lawyer(number + 100)
        _process(lawyer, civil, f"FIFTY-{number}", REPORT_START)

    fifty_lawyers_queries = _workload_query_count()

    assert one_lawyer_queries == fifty_lawyers_queries
    assert fifty_lawyers_queries <= MAX_LAWYERS_WORKLOAD_QUERIES == 6


def test_workload_export_aggregates_a_fallo_process_once(monkeypatch):
    """Fails if a multi-stage failed process is duplicated in the workload totals."""
    monkeypatch.setattr("gym_app.views.reports.user_id", None)
    civil = Case.objects.create(type="Civil")
    labor = Case.objects.create(type="Laboral")
    primary = _lawyer(1)
    secondary = _lawyer(2)
    fallo = Stage.objects.create(status="Fallo")
    review = Stage.objects.create(status="Revisión")
    _process(primary, civil, "ACTIVE", REPORT_START)
    _process(primary, civil, "FAILED-TWICE", REPORT_START, (review, fallo))
    _process(primary, labor, "FAILED", REPORT_START, (fallo,))
    _process(secondary, civil, "SECOND", REPORT_START)

    response = generate_lawyers_workload_report(_response(), REPORT_START, REPORT_END)
    dataframe = pd.read_excel(
        io.BytesIO(response.content), sheet_name="Carga de Trabajo de Abogados"
    )
    primary_row = dataframe.loc[dataframe["Email"] == primary.email].iloc[0]

    assert primary_row["Total de Procesos Asignados"] == 3
    assert primary_row["Procesos Activos"] == 1
    assert primary_row["Procesos Completados"] == 2
    assert primary_row["Distribución por Tipo de Caso"] == "Civil: 2, Laboral: 1"


def test_workload_export_orders_lawyers_by_total_processes(monkeypatch):
    """Fails if the Excel export stops placing the highest workload first."""
    monkeypatch.setattr("gym_app.views.reports.user_id", None)
    civil = Case.objects.create(type="Civil")
    primary = _lawyer(1)
    secondary = _lawyer(2)
    _process(primary, civil, "FIRST-ONE", REPORT_START)
    _process(primary, civil, "FIRST-TWO", REPORT_START)
    _process(secondary, civil, "SECOND", REPORT_START)

    response = generate_lawyers_workload_report(_response(), REPORT_START, REPORT_END)
    dataframe = pd.read_excel(io.BytesIO(response.content))

    assert dataframe.iloc[0]["Email"] == primary.email


def test_workload_export_writes_charts_to_its_chart_worksheet(monkeypatch):
    """Fails if a multi-lawyer export stops including its two workload charts."""
    monkeypatch.setattr("gym_app.views.reports.user_id", None)
    civil = Case.objects.create(type="Civil")
    _process(_lawyer(1), civil, "FIRST", REPORT_START)
    _process(_lawyer(2), civil, "SECOND", REPORT_START)

    response = generate_lawyers_workload_report(_response(), REPORT_START, REPORT_END)
    workbook = load_workbook(io.BytesIO(response.content))

    assert "Carga de Trabajo de Abogados" in workbook.sheetnames
    assert "Gráficos" in workbook.sheetnames
    assert len(workbook["Gráficos"]._charts) == 2


def test_workload_export_selects_the_requested_lawyer(monkeypatch):
    """Fails if the report ignores the lawyer selected by the report control."""
    civil = Case.objects.create(type="Civil")
    selected = _lawyer(1)
    other = _lawyer(2)
    _process(selected, civil, "START", REPORT_START)
    _process(other, civil, "OTHER", REPORT_START)
    monkeypatch.setattr("gym_app.views.reports.user_id", selected.pk)

    response = generate_lawyers_workload_report(_response(), REPORT_START, REPORT_END)
    dataframe = pd.read_excel(io.BytesIO(response.content))

    assert dataframe["Email"].tolist() == [selected.email]
    assert dataframe.iloc[0]["Total de Procesos Asignados"] == 1


def test_workload_export_includes_processes_on_both_date_boundaries(monkeypatch):
    """Fails if the created-at range excludes either endpoint of the selected period."""
    monkeypatch.setattr("gym_app.views.reports.user_id", None)
    civil = Case.objects.create(type="Civil")
    lawyer = _lawyer(1)
    _process(lawyer, civil, "START", REPORT_START)
    _process(lawyer, civil, "END", REPORT_END)
    _process(lawyer, civil, "OUTSIDE", REPORT_END + datetime.timedelta(microseconds=1))

    response = generate_lawyers_workload_report(_response(), REPORT_START, REPORT_END)
    dataframe = pd.read_excel(io.BytesIO(response.content))

    assert dataframe.iloc[0]["Total de Procesos Asignados"] == 2


def test_workload_export_writes_the_workload_sheet_when_no_processes_match(monkeypatch):
    """Fails if an empty report period stops producing the workload worksheet."""
    monkeypatch.setattr("gym_app.views.reports.user_id", None)

    response = generate_lawyers_workload_report(_response(), REPORT_START, REPORT_END)
    dataframe = pd.read_excel(io.BytesIO(response.content))
    workbook = load_workbook(io.BytesIO(response.content))

    assert len(dataframe.index) == 0
    assert workbook.sheetnames == ["Carga de Trabajo de Abogados"]
